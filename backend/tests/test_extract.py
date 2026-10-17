"""Tests for app.extract.extract_text (plan §6.3 step 5, §6.5 step 2, decision D23).

Every sample document is built in memory. No fixture file is read from disk.
- DOCX samples come from python-docx's own Document() writer.
- PDF samples come from pypdf.PdfWriter (blank pages), or from a minimal
  hand-built PDF byte string with one text line for the happy path.
- TXT samples are plain bytes.

extract.py does no I/O and reads no settings, so these are plain unit tests.
No integration marker and no environment setup are needed.
"""

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.extract import UnsupportedDocument, _normalize_pdf_page, extract_text


def _docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _blank_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _text_pdf_bytes(text: str) -> bytes:
    """Build a minimal one-page PDF whose page content is one Helvetica text line.

    The xref offsets are computed from the actual object bytes, so the file is
    valid without relying on pypdf's xref reconstruction.
    """
    stream = f"BT /F1 24 Tf 20 100 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def _one_char_per_line_pdf_bytes(text: str) -> bytes:
    """Build a PDF whose content stream positions every character on its own line.

    Reproduces the real pypdf extraction artifact seen on some generated PDFs
    (e.g. "Employment at Anthropic.pdf"): a Td (move text position) before every
    single Tj (show text) call, each moving down one line, makes pypdf's own
    line-break heuristic insert "\\n" between every character.
    """
    ops = ["BT /F1 24 Tf 20 190 Td"]
    for char in text:
        escaped = char.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({escaped}) Tj 0 -10 Td")
    ops.append("ET")
    stream = " ".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 400] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


# TXT


def test_txt_returns_decoded_text_exactly() -> None:
    assert extract_text(b"Hello plain text.", "txt") == "Hello plain text."


@pytest.mark.parametrize("extension", ["TXT", ".txt", ".TXT"])
def test_txt_extension_is_case_insensitive_and_dot_tolerant(extension: str) -> None:
    assert extract_text(b"same content", extension) == "same content"


def test_txt_malformed_utf8_becomes_replacement_characters_without_raising() -> None:
    content = b"valid start \xff\xfe valid end"
    result = extract_text(content, "txt")
    assert "�" in result
    assert result.startswith("valid start ")
    assert result.endswith(" valid end")


def test_txt_empty_bytes_raise_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(b"", "txt")


def test_txt_whitespace_only_bytes_raise_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(b"   \n\t  ", "txt")


# PDF


def test_pdf_with_real_text_returns_the_page_text() -> None:
    result = extract_text(_text_pdf_bytes("Hello PDF"), "pdf")
    assert result.strip() == "Hello PDF"


def test_pdf_bytes_without_pdf_magic_raise_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(b"not a pdf at all", "pdf")


@pytest.mark.parametrize("extension", ["PDF", ".pdf", ".PDF"])
def test_pdf_wrong_magic_is_rejected_for_every_extension_variant(extension: str) -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(b"not a pdf at all", extension)


def test_pdf_with_magic_but_truncated_body_raises_unsupported_document() -> None:
    truncated = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog"
    with pytest.raises(UnsupportedDocument):
        extract_text(truncated, "pdf")


def test_pdf_with_no_text_content_raises_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(_blank_pdf_bytes(), "pdf")


def test_pdf_one_char_per_line_artifact_is_reconstructed() -> None:
    """Regression test: a PDF that triggers pypdf's one-glyph-per-line extraction
    must still come out as normal prose, not exploded into one-word "paragraphs"
    by chunk_text()'s blank-line paragraph detection.
    """
    result = extract_text(_one_char_per_line_pdf_bytes("agent hours outnumber"), "pdf")
    assert "agent hours outnumber" in result.replace("\n\n", "")


def test_normalize_pdf_page_collapses_a_degenerate_page() -> None:
    degenerate = "H\no\nw\n \na\nr\ne\n \ny\no\nu\n"
    assert _normalize_pdf_page(degenerate) == "How are you"


def test_normalize_pdf_page_leaves_a_normal_page_unchanged() -> None:
    normal = "This is a normal line.\nThis is another wrapped line.\n\nNew paragraph."
    assert _normalize_pdf_page(normal) == normal


def test_normalize_pdf_page_handles_empty_text() -> None:
    assert _normalize_pdf_page("") == ""


# DOCX


def test_docx_paragraphs_are_joined_with_newline_exactly() -> None:
    content = _docx_bytes("Hello world", "Second paragraph")
    assert extract_text(content, "docx") == "Hello world\nSecond paragraph"


@pytest.mark.parametrize("extension", ["DOCX", ".docx", ".DOCX"])
def test_docx_extension_is_case_insensitive_and_dot_tolerant(extension: str) -> None:
    assert extract_text(_docx_bytes("Hello world"), extension) == "Hello world"


def test_docx_with_no_paragraphs_raises_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(_docx_bytes(), "docx")


def test_docx_with_only_empty_paragraphs_raises_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(_docx_bytes("", "   "), "docx")


def test_docx_bytes_that_are_not_a_zip_raise_unsupported_document() -> None:
    with pytest.raises(UnsupportedDocument):
        extract_text(b"not a docx file", "docx")


# Cross-cutting


@pytest.mark.parametrize("extension", ["xlsx", "doc", ""])
def test_unsupported_extension_raises_and_names_the_extension(extension: str) -> None:
    with pytest.raises(UnsupportedDocument, match="unsupported extension"):
        extract_text(b"Some readable text", extension)


def test_legacy_doc_extension_is_rejected_even_with_valid_docx_bytes() -> None:
    with pytest.raises(UnsupportedDocument, match="unsupported extension: 'doc'"):
        extract_text(_docx_bytes("Hello world"), "doc")


def test_unsupported_document_is_an_exception_subclass() -> None:
    assert issubclass(UnsupportedDocument, Exception)
