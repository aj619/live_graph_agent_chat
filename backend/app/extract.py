"""Extract plain text from uploaded PDF, TXT and DOCX documents (plan §6.3 step 5, §6.5 step 2).

The output is one plain string. Callers pass it to chunking.chunk_text().

Supported extensions (plan decision D23: DOCX only, no .doc, no LibreOffice):
- "pdf": the bytes must start with b"%PDF" (checked from byte 0). pypdf extracts the
  text of every page; pages are joined with a blank line. A pypdf artifact on some
  PDFs (every glyph on its own line) is normalized away first (_normalize_pdf_page),
  since it otherwise makes chunking.chunk_text() misread every inter-word space as
  a paragraph break.
- "docx": the bytes must open with python-docx. The text of every paragraph is joined
  with a newline.
- "txt": the bytes are decoded as UTF-8 with replacement, so malformed bytes become
  U+FFFD and never raise.

Any other extension is rejected. Extension matching is case-insensitive, and a
leading dot is allowed. Route-level checks such as ALLOWED_EXTENSIONS and the
file-size limit belong to the upload route, not to this module.

A document that fails its content check, that a parser cannot open, or whose
extracted text is empty or whitespace-only raises UnsupportedDocument.

The function is pure: it works on in-memory bytes, does no I/O, does no logging,
and reads no settings.
"""

import io

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

_PDF_MAGIC = b"%PDF"
_PDF_PAGE_SEPARATOR = "\n\n"
_DOCX_PARAGRAPH_SEPARATOR = "\n"


class UnsupportedDocument(Exception):
    """Raised when the bytes are not a supported document or yield no text."""


def extract_text(content: bytes, extension: str) -> str:
    """Return the text of content, a document with the given file extension.

    extension is the file name's extension, with or without a leading dot. Matching
    is case-insensitive. Raises UnsupportedDocument if the extension is not pdf, txt
    or docx, if the bytes do not match the extension, or if no text is extracted.
    """
    kind = extension.lower().removeprefix(".")
    if kind == "pdf":
        text = _extract_pdf(content)
    elif kind == "docx":
        text = _extract_docx(content)
    elif kind == "txt":
        text = content.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocument(f"unsupported extension: {extension!r}")

    if not text.strip():
        raise UnsupportedDocument(f"no text could be extracted from the {kind} document")
    return text


def _extract_pdf(content: bytes) -> str:
    if not content.startswith(_PDF_MAGIC):
        raise UnsupportedDocument("PDF content does not start with %PDF")
    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [_normalize_pdf_page(page.extract_text() or "") for page in reader.pages]
    except (PdfReadError, ValueError, TypeError, KeyError, IndexError, OSError) as exc:
        raise UnsupportedDocument(f"PDF could not be parsed: {exc}") from exc
    return _PDF_PAGE_SEPARATOR.join(pages)


def _normalize_pdf_page(text: str) -> str:
    """Undo a pypdf extraction artifact seen on some PDFs: every glyph (including
    plain spaces between words) lands on its own line, so most "lines" are zero or
    one characters long. Left as is, chunking.chunk_text()'s paragraph-boundary
    regex (a blank line, "\\n\\s*\\n") matches every inter-word space, exploding the
    whole page into one-word "paragraphs" and badly skewing where chunk boundaries
    fall. If most lines in this page are that short, the newlines carry no real
    line-wrap information, so they are dropped before chunking ever sees them.
    Unaffected (normally wrapped) pages are returned unchanged.
    """
    lines = text.split("\n")
    if not lines:
        return text
    short_lines = sum(1 for line in lines if len(line) <= 1)
    if short_lines / len(lines) > 0.5:
        return "".join(lines)
    return text


def _extract_docx(content: bytes) -> str:
    try:
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
    except Exception as exc:
        raise UnsupportedDocument(f"DOCX could not be parsed: {exc}") from exc
    return _DOCX_PARAGRAPH_SEPARATOR.join(paragraphs)
